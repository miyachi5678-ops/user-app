"""解析結果をExcelファイルに出力する"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from .order_parser import Order
from .order_processor import OrderDetail


# 色定義
COLOR_HEADER_BG  = "1F4E79"   # 濃紺
COLOR_HEADER_FG  = "FFFFFF"   # 白
COLOR_SUBHDR_BG  = "BDD7EE"   # 薄青
COLOR_UNKNOWN_BG = "FFE0E0"   # 薄赤（構成未登録）
COLOR_ALT_BG     = "F2F2F2"   # 薄グレー（交互）


def write_report(
    orders: list[Order],
    details: list[OrderDetail],
    unknown_hinban: list[str],
    output_path: str,
):
    """3シート構成のExcelレポートを生成する"""
    wb = openpyxl.Workbook()

    _write_summary_sheet(wb, orders, details, unknown_hinban)
    _write_detail_sheet(wb, details)
    _write_order_sheet(wb, orders)

    # デフォルトシートを削除
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    wb.save(output_path)
    return output_path


def _write_summary_sheet(wb, orders, details, unknown_hinban):
    """シート1: 集計（子品番・長さ別の合計必要数）"""
    from .order_processor import summarize_by_length

    ws = wb.create_sheet("集計")
    generated_at = datetime.now().strftime("%Y/%m/%d %H:%M")

    ws["A1"] = f"注文解析レポート　生成日時: {generated_at}"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = f"対象注文件数: {len(orders)}件　　構成未登録品番: {len(unknown_hinban)}件"
    ws["A2"].font = Font(color="FF0000" if unknown_hinban else "000000")

    if unknown_hinban:
        ws["A3"] = "【未登録品番】" + "　".join(unknown_hinban)
        ws["A3"].font = Font(color="FF0000")

    headers = ["子品番", "長さ(mm)", "形状", "合計数量", "該当納期"]
    start_row = 5
    _write_header_row(ws, start_row, headers)

    summary = summarize_by_length(details)
    for i, row in enumerate(summary):
        r = start_row + 1 + i
        fill = PatternFill("solid", fgColor=COLOR_ALT_BG) if i % 2 == 0 else None
        is_unknown = row["子品番"] is None

        values = [
            row["子品番"] or "（未登録）",
            row["直_長さ"],
            row["形状ラベル"] or row["形状"],
            row["合計数量"],
            row["納期一覧"],
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=val)
            if is_unknown:
                cell.fill = PatternFill("solid", fgColor=COLOR_UNKNOWN_BG)
            elif fill:
                cell.fill = fill
            cell.alignment = Alignment(horizontal="left")

    _set_column_widths(ws, [20, 12, 30, 12, 40])


def _write_detail_sheet(wb, details):
    """シート2: 注文明細（オーダー別・子品番別）"""
    ws = wb.create_sheet("注文明細")

    headers = ["オーダーNo", "親品番", "子品番", "数量", "最新納期", "工場", "形状", "R側", "長さ(mm)", "L側"]
    _write_header_row(ws, 1, headers)

    for i, d in enumerate(details):
        r = 2 + i
        fill = PatternFill("solid", fgColor=COLOR_ALT_BG) if i % 2 == 0 else None
        is_unknown = d.子品番 is None

        values = [
            d.オーダーNo, d.親品番, d.子品番 or "（未登録）",
            d.数量, d.最新納期, d.工場,
            d.形状ラベル or d.形状, d.R側, d.直_長さ, d.L側,
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=val)
            if is_unknown:
                cell.fill = PatternFill("solid", fgColor=COLOR_UNKNOWN_BG)
            elif fill:
                cell.fill = fill

    _set_column_widths(ws, [18, 16, 16, 8, 14, 8, 28, 8, 10, 8])


def _write_order_sheet(wb, orders):
    """シート3: 注文一覧（③ジュケンの内容そのまま）"""
    ws = wb.create_sheet("注文一覧")

    headers = ["オーダーNo", "品番", "発注数", "最新納期", "工場", "発注日"]
    _write_header_row(ws, 1, headers)

    for i, o in enumerate(orders):
        r = 2 + i
        fill = PatternFill("solid", fgColor=COLOR_ALT_BG) if i % 2 == 0 else None
        values = [o.オーダーNo, o.品番, o.発注数, o.最新納期, o.工場, o.発注日]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=val)
            if fill:
                cell.fill = fill

    _set_column_widths(ws, [18, 16, 8, 14, 8, 14])


def _write_header_row(ws, row: int, headers: list[str]):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = Font(bold=True, color=COLOR_HEADER_FG)
        cell.fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
        cell.alignment = Alignment(horizontal="center")


def _set_column_widths(ws, widths: list[float]):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
