"""③ジュケンシートから注文データを読み込む"""
import openpyxl
from dataclasses import dataclass
from datetime import datetime, date


@dataclass
class Order:
    オーダーNo: str
    品番: str
    発注数: float
    納期: str | None
    最新納期: str | None
    工場: str | None
    発注日: str | None


def parse_order_excel(excel_path: str) -> list[Order]:
    """③ジュケンシートから注文リストを返す"""
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    sheet_name = None
    for name in wb.sheetnames:
        if "ジュケン" in name or "juken" in name.lower():
            sheet_name = name
            break
    if sheet_name is None:
        raise ValueError(f"ジュケンシートが見つかりません。シート一覧: {wb.sheetnames}")

    ws = wb[sheet_name]
    orders = []

    # ヘッダー行を動的に検出して次の行からデータ読み込み
    # A列=納期, B列=最新納期, C列=品番, D列=オーダーNo, E列=発注数, F列=納期回答, G列=工場, H列=発注日
    data_start_row = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        if row[3] and "オーダー" in str(row[3]):
            data_start_row = i + 1
            break
    if data_start_row is None:
        data_start_row = 4  # フォールバック

    for row in ws.iter_rows(min_row=data_start_row, values_only=True):
        # A列=納期, B列=最新納期, C列=品番, D列=オーダーNo, E列=発注数, F列=納期回答, G列=工場, H列=発注日
        noki, saishinnoki, hinban, order_no, hatchusuu, _, kojyo, hatchubi, *_ = (
            list(row) + [None] * 15
        )[:15]

        if not hinban or not order_no:
            continue
        if str(order_no).strip() in ("計画", ""):
            continue

        # 発注数が数値でない行（ヘッダーの混入など）はスキップ
        try:
            hatchu_float = float(hatchusuu) if hatchusuu else 1.0
        except (ValueError, TypeError):
            continue

        orders.append(Order(
            オーダーNo=str(order_no).strip(),
            品番=str(hinban).strip(),
            発注数=hatchu_float,
            納期=_format_date(noki),
            最新納期=_format_date(saishinnoki),
            工場=str(kojyo).strip() if kojyo else None,
            発注日=_format_date(hatchubi),
        ))



    return orders


def _format_date(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y/%m/%d")
    s = str(value).strip()
    # シリアル値（数値）が来た場合はスキップ
    try:
        float(s)
        return None
    except ValueError:
        return s if s else None
